"""DOCX uploads preserve paragraph/table order, nested tables, and merged cells."""
import asyncio
import io
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # apps/api

from docx import Document  # noqa: E402
from starlette.datastructures import UploadFile  # noqa: E402

from app.api.v1 import knowledge  # noqa: E402
from app.api.v1.knowledge import MaterialIn  # noqa: E402


def test_docx_upload_extracts_body_order_nested_tables_and_merged_cells():
    document = Document()
    document.add_paragraph("段落一")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "合并单元格"
    table.cell(0, 1).text = "右上单元格"
    table.cell(1, 1).text = "右下单元格"
    table.cell(0, 0).merge(table.cell(1, 0))
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套表格"
    document.add_paragraph("段落二")

    raw = io.BytesIO()
    document.save(raw)
    result = asyncio.run(
        knowledge.upload_material(
            UploadFile(file=io.BytesIO(raw.getvalue()), filename=f"顺序-{uuid.uuid4().hex}.docx")
        )
    )

    assert result["ok"] is True
    uploaded = knowledge._load()[0]["content"]
    assert uploaded == "段落一\n合并单元格\n嵌套表格\n右上单元格\n右下单元格\n段落二"
    assert uploaded.count("合并单元格") == 1


def test_material_metadata_fields_round_trip_through_existing_store():
    metadata = {
        "context_sources": [{"kind": "asset", "id": "m_source", "excerpt": "资料选段"}],
        "source_material_ids": ["m_source"],
        "claim_check": {"status": "needs_review", "claims": []},
    }
    result = asyncio.run(
        knowledge.add_material(
            MaterialIn(title="带来源草稿", content="这是一段超过十个字符的测试正文。", **metadata)
        )
    )

    item = next(row for row in knowledge._load() if row["title"] == "带来源草稿")
    assert result["data"]["item"] == item
    assert {key: item[key] for key in metadata} == metadata
